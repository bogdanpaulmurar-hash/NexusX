# -*- coding: utf-8 -*-
"""Engine tests (stdlib unittest; no network except a local test server)."""
import http.server
import json
import shutil
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from nexusx.codex import exporter, linkmap, summary  # noqa: E402
from nexusx.codex.redact import redact  # noqa: E402
from nexusx.codex.scanner import ScanOptions, enrich, scan_folder  # noqa: E402
from nexusx.links import checker, reporter  # noqa: E402
from nexusx.links.extractor import extract_from_file  # noqa: E402
from nexusx.links.models import LinkRecord, status_group  # noqa: E402
from nexusx import mailer  # noqa: E402


class CodeExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        a = self.tmp / "ProjA"
        (a / "pkg_a").mkdir(parents=True)
        (a / "pkg_a" / "__init__.py").write_text("")
        (a / "pkg_a" / "core.py").write_text(
            'import os\nimport requests\nURL = "https://siemens.sharepoint.com/teams/X/page.aspx"\n'
            'PASSWORD = "SuperSecret123"\nDB = "Server=sql01;Database=x;User Id=u;Password=hunter2;"\n'
            'OUT = r"C:\\Users\\me\\ProjB\\data"\n', encoding="utf-8")
        (a / "node_modules").mkdir()
        (a / "node_modules" / "x.js").write_text("x")
        (a / ".hidden.py").write_text("x")
        (a / "big.bin").write_bytes(b"\x00\x01" * 100)
        b = self.tmp / "ProjB"
        (b / "pkg_b").mkdir(parents=True)
        (b / "pkg_b" / "__init__.py").write_text("")
        (b / "pkg_b" / "main.py").write_text(
            "from pkg_a.core import URL\nimport webbrowser\nwebbrowser.open('https://siemens.sharepoint.com/teams/X/page.aspx')\n"
            "for i in range(3000):\n    pass\n" + "\n".join(f"# line {i} " + "x" * 80 for i in range(400)))
        self.a, self.b = a, b

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_scan_excludes(self):
        files = scan_folder(self.a, ScanOptions.full())
        names = {f.name for f in files}
        self.assertIn("core.py", names)
        self.assertNotIn("x.js", names)
        self.assertNotIn(".hidden.py", names)
        infos = enrich(files, self.a)
        self.assertTrue(any(i.binary for i in infos if i.path.name == "big.bin"))

    def test_redaction(self):
        txt, n = redact('PASSWORD = "SuperSecret123"\nconn="Server=a;Password=hunter2;"\nAuthorization: Bearer abcdefghijkl')
        self.assertNotIn("SuperSecret123", txt)
        self.assertNotIn("hunter2", txt)
        self.assertNotIn("abcdefghijkl", txt)
        self.assertGreaterEqual(n, 3)

    def test_export_all_formats_and_split(self):
        files = scan_folder(self.tmp, ScanOptions.full())
        for fmt in exporter.FORMATS:
            res = exporter.export_files(self.tmp / "out" / f"exp{fmt}", fmt, files, self.tmp, "T",
                                        exporter.ExportSettings(line_numbers=True))
            self.assertEqual(len(res.parts), 1)
            content = res.parts[0].read_text(encoding="utf-8")
            self.assertNotIn("SuperSecret123", content)
            self.assertTrue(any(x.endswith("big.bin") for x in res.skipped_binary))
            if fmt == ".json":
                json.loads(content)
        res = exporter.export_files(self.tmp / "out" / "split.txt", ".txt", files, self.tmp, "T",
                                    exporter.ExportSettings(split_bytes=20_000))
        self.assertGreater(len(res.parts), 1)
        first = res.parts[0].read_text(encoding="utf-8")
        self.assertIn(f"Parte    : 1/{len(res.parts)}", first)
        self.assertIn("CONTINUA CU", first)
        self.assertIn("EXPORT COMPLET", res.parts[-1].read_text(encoding="utf-8"))
        res = exporter.export_files(self.tmp / "out" / "split.json", ".json", files, self.tmp, "T",
                                    exporter.ExportSettings(split_bytes=20_000))
        self.assertGreater(len(res.parts), 1)
        for p in res.parts:
            json.loads(p.read_text(encoding="utf-8"))
        z = exporter.zip_parts(res.parts, self.tmp / "out" / "x.zip")
        self.assertEqual(len(zipfile.ZipFile(z).namelist()), len(res.parts))

    def test_summary(self):
        for fmt in (".md", ".txt", ".html", ".json"):
            p = summary.build_summary(self.a, self.tmp / f"sum{fmt}", fmt, {"binary": True})
            self.assertTrue(p.exists())
            self.assertNotIn("SuperSecret123", p.read_text(encoding="utf-8"))

    def test_linkmap_cross_project(self):
        res = linkmap.analyse([("ProjA", self.a), ("ProjB", self.b)])
        kinds = {e.kind for e in res.edges}
        self.assertTrue({"url", "import", "db", "path", "automation"} <= kinds, kinds)
        links = res.project_links()
        self.assertIn(("ProjB", "ProjA"), links)
        shared = res.shared_resources()
        self.assertTrue(any(k == "url" and len(p) == 2 for k, _t, p in shared))
        self.assertFalse(any("hunter2" in e.target for e in res.edges))
        for fn in (linkmap.to_markdown, linkmap.to_json, linkmap.to_csv, linkmap.to_html):
            self.assertTrue(fn(res))
        imp = [e for e in res.edges if e.kind == "import" and e.target == "os"]
        self.assertEqual(imp[0].scope, "stdlib")


def _make_docx(path: Path):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    h = doc.add_heading("Kapitel Eins", level=1)
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "1"); bs.set(qn("w:name"), "Ziel1")
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "1")
    h._p.append(bs); h._p.append(be)
    p = doc.add_paragraph("Visit https://example.org/text-url now. ")
    rid = doc.part.relate_to("https://example.com/hyper", RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "Hyper"; r.append(t); hl.append(r); p._p.append(hl)
    doc.add_heading("Kapitel Zwei", level=1)
    p2 = doc.add_paragraph("See ")
    for anchor in ("Ziel1", "Fehlt"):
        hl2 = OxmlElement("w:hyperlink"); hl2.set(qn("w:anchor"), anchor)
        r2 = OxmlElement("w:r"); t2 = OxmlElement("w:t"); t2.text = anchor; r2.append(t2); hl2.append(r2); p2._p.append(hl2)

    def field(par, instr, result):
        for typ, txt in (("begin", None), ("instr", instr), ("separate", None), ("res", result), ("end", None)):
            run = OxmlElement("w:r")
            if typ in ("begin", "separate", "end"):
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); run.append(fc)
            elif typ == "instr":
                it = OxmlElement("w:instrText"); it.text = txt; run.append(it)
            else:
                tt = OxmlElement("w:t"); tt.text = txt; run.append(tt)
            par._p.append(run)
    p3 = doc.add_paragraph("Fields: ")
    field(p3, ' HYPERLINK "https://example.net/field" ', "field link")
    field(p3, " REF Ziel1 \\h ", "Kapitel Eins")
    field(p3, " PAGEREF Verloren \\h ", "3")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 1).text = "cell http://example.org/in-table"
    doc.sections[0].header.paragraphs[0].text = "Header link https://example.org/header"
    doc.save(path)


class ExtractorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_docx(self):
        p = self.tmp / "t.docx"
        _make_docx(p)
        recs = {r.url: r for r in extract_from_file(p)}
        for u in ("https://example.org/text-url", "https://example.com/hyper", "https://example.net/field",
                  "http://example.org/in-table", "https://example.org/header"):
            self.assertIn(u, recs, list(recs))
        self.assertEqual(recs["#bookmark=Ziel1"].status, "Bookmark OK")
        self.assertIn("Kapitel Eins", recs["#bookmark=Ziel1"].target_section)
        self.assertEqual(recs["#bookmark=Fehlt"].status, "Broken Bookmark")
        self.assertEqual(recs["#bookmark=Verloren"].status, "Broken Bookmark")
        self.assertIn("Table 1, Row 1, Col 2", recs["http://example.org/in-table"].location)
        self.assertTrue(recs["https://example.org/header"].location.startswith("Header"))
        self.assertEqual(recs["https://example.com/hyper"].section, "Kapitel Eins")
        self.assertIn("cross-reference", recs["#bookmark=Ziel1"].source)

    def test_md_html(self):
        md = self.tmp / "a.md"
        md.write_text("# Intro\n[ok](#intro) [bad](#nope) [w](https://x.org/a_(b)) <https://y.org>\n[f](other.md)\n", encoding="utf-8")
        r = {x.url: x for x in extract_from_file(md)}
        self.assertEqual(r["#bookmark=intro"].status, "Bookmark OK")
        self.assertEqual(r["#bookmark=nope"].status, "Broken Bookmark")
        self.assertIn("https://x.org/a_(b)", r)
        self.assertIn("https://y.org", r)
        self.assertTrue(any(u.startswith("file:") for u in r))
        h = self.tmp / "a.html"
        h.write_text('<html><body><h1 id="top">T</h1><a href="#top">t</a><a href="#x">x</a>'
                     '<a href="https://z.org/p">z</a><a href="mailto:a@b.de">m</a><img src="img.png"></body></html>')
        r = {x.url: x for x in extract_from_file(h)}
        self.assertEqual(r["#bookmark=top"].status, "Bookmark OK")
        self.assertEqual(r["#bookmark=x"].status, "Broken Bookmark")
        self.assertIn("https://z.org/p", r)
        self.assertIn("mailto:a@b.de", r)


class _Handler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        p = self.path
        if p == "/ok":
            self._send(200, '<html><head><title>Good page</title></head><body><h1>Welcome</h1><div id="sec">x</div>' + "text " * 600 + "</body></html>")
        elif p == "/soft404":
            self._send(200, "<html><head><title>Page not found</title></head><body>Sorry</body></html>")
        elif p == "/long-forbidden":
            self._send(200, "<html><head><title>Docs</title></head><body>" + "The word forbidden appears in a long article. " * 100 + "</body></html>")
        elif p == "/login":
            self.send_response(302); self.send_header("Location", "/adfs/ls/?wa=x"); self.end_headers()
        elif p.startswith("/adfs/"):
            self._send(200, "<html><head><title>Sign In</title></head><body>login</body></html>")
        elif p == "/dl":
            self.send_response(200); self.send_header("Content-Type", "application/octet-stream")
            self.send_header("Content-Disposition", 'attachment; filename="report.xlsx"'); self.send_header("Content-Length", "4096"); self.end_headers()
            self.wfile.write(b"x" * 4096)
        elif p == "/forbidden":
            self._send(403, "no")
        else:
            self._send(404, "missing")

    def _send(self, code, body):
        data = body.encode()
        self.send_response(code)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


class CheckerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()

    def test_statuses(self):
        cases = {"/ok": "OK", "/soft404": "Page error", "/long-forbidden": "OK", "/login": "Authentication required",
                 "/dl": "OK (Download)", "/forbidden": "Access denied", "/nothere": "Broken",
                 "/ok#sec": "OK", "/ok#missing": "Anchor missing"}
        recs = [LinkRecord(file="f", path="p", location="l", url=self.base + k) for k in cases]
        recs.append(LinkRecord(file="f", path="p", location="l", url=self.base + "/ok"))  # duplicate → cache
        recs.append(LinkRecord(file="f", path="p", location="l", url="mailto:someone@siemens.com"))
        recs.append(LinkRecord(file="f", path="p", location="l", url="file:///definitely/not/here.docx"))
        checker.check_many(recs, timeout=5, workers=4)
        for r, (k, expected) in zip(recs, cases.items()):
            self.assertEqual(r.status, expected, f"{k}: {r.status} {r.error}")
        self.assertEqual(recs[len(cases)].status, "OK")
        self.assertEqual(recs[-2].status, "Mailto (syntax OK)")
        self.assertEqual(recs[-1].status, "Broken (path not found)")
        self.assertEqual(status_group("Authentication required"), "auth")

    def test_crawl_and_reports(self):
        from nexusx.links.crawler import expand_page
        out = expand_page(self.base + "/ok", timeout=5)
        self.assertEqual(out, [])  # page has no links
        tmp = Path(tempfile.mkdtemp())
        recs = [LinkRecord(file="doc.docx", path="p", location="Paragraph 1", url="https://a.org", status="OK"),
                LinkRecord(file="doc.docx", path="p", location="Paragraph 2", url="https://b.org", status="Broken", notes="fix")]
        paths = reporter.export_all(recs, tmp, formats=(".xlsx", ".docx", ".csv", ".md", ".json", ".html"))
        self.assertEqual(len(paths), 6)
        html_body = mailer.build_links_html(recs, "de", "Bogdan")
        self.assertIn("Fehlerhaft: 1", html_body)
        ok, msg = mailer._eml("S", html_body, paths, "a@b.c", "", tmp, "")
        self.assertTrue(ok)
        self.assertTrue(list(tmp.glob("*.eml")))
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
